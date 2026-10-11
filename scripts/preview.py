"""Disposable synthetic preview. Uses a separate database, never the real instance data."""
import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from sqlalchemy import select
from werkzeug.security import generate_password_hash
from app import create_app
from app.models import User, Job, Candidate, Resume, Evaluation, Recruitment, StageEvent, now
from app.db import transaction
from app.ai import validate_report

app = create_app({"DATA_DIR": str(Path(__file__).resolve().parents[1] / "test-results" / "preview"),
                  "START_WORKER": False, "DEMO_MODE": True})

with transaction(app, write=True) as db:
    if not db.scalar(select(User.id)):
        db.add(User(username="demo", password_hash=generate_password_hash("Preview-only-2026")))
        jobs = []
        for name, skill in [("Python 后端开发实习生", "Python"), ("数据分析实习生", "SQL"), ("前端开发实习生", "JavaScript")]:
            criteria = [{"key": "skill", "type": "must", "name": f"{skill}基础", "requirement": f"能够使用{skill}完成项目", "max_score": 60,
                         "scoring_rule": "项目证据完整60，部分实践30，未提供0"},
                        {"key": "internship", "type": "must", "name": "相关实习经验", "requirement": "有至少一次岗位相关实习", "max_score": 25,
                         "scoring_rule": "相关实习25，非相关实习10，未提供0"},
                        {"key": "project", "type": "bonus", "name": "完整项目实践", "requirement": "具有完整项目交付经历", "max_score": 15,
                         "scoring_rule": "完整实践15，部分实践8，未提供0"}]
            j = Job(name=name, jd=f"寻找熟悉{skill}的实习生，有项目和实习经历。", description="上海 · 技术团队 · 实习", criteria=criteria)
            db.add(j)
            db.flush()
            jobs.append(j)
        stages = ["group_interview", "technical_1", "technical_2", "hr", "offer_signed", "onboarded", "rejected", "pending_interview"]
        for i, name in enumerate(["林予安", "陈知夏", "周亦舟", "许沐晴", "沈嘉禾", "陆星言", "宋书宁", "顾清越", "江景行", "苏念"]):
            j = jobs[i % 3]
            c = Candidate(name="示例·" + name, education="本科", major="计算机科学", email=f"example{i}@example.test")
            db.add(c)
            db.flush()
            text = "完成岗位相关项目，负责接口开发、数据处理和交付。\n有一段相关实习经历。"
            r = Resume(candidate_id=c.id, filename="合成示例简历.docx", path="example.docx", sha256=str(i), markdown=text, status="succeeded")
            db.add(r)
            db.flush()
            snapshot = {"candidate": {"name": c.name, "education": c.education, "major": c.major}, "job_name": j.name, "jd": j.jd,
                        "criteria": j.criteria, "criteria_version": 1, "markdown": text, "resume_revision": 1}
            raw = {"relevant_information": {"skills": ["岗位技能实践"], "internships": ["相关实习经历（合成示例）"], "projects": ["岗位相关项目（合成示例）"]},
                   "criteria_results": [{"criterion_key": row["key"], "status": "partial" if i % 4 == 0 and row["key"] == "internship" else "met",
                                         "score": max(0, row["max_score"] - (i % 3) * 3 - (10 if i % 4 == 0 and row["key"] == "internship" else 0)),
                                         "reason": "合成示例：已有实践依据，具体深度需面试核实。",
                                         "evidence": [{"quote": "有一段相关实习经历。" if row["key"] == "internship" else "完成岗位相关项目", "location": "合成示例"}]} for row in j.criteria],
                   "strengths": ["有项目与岗位相关实践"], "gaps": ["需核实独立完成程度"], "interview_questions": ["请介绍项目职责与遇到的技术难点。"]}
            report = validate_report(raw, snapshot)
            ev = Evaluation(candidate_id=c.id, resume_id=r.id, job_id=j.id, status="succeeded", snapshot=snapshot,
                            report=report, total_score=report["total_score"], model="合成示例（未调用模型）", completed_at=now(),
                            batch_key=f"preview-{i}", decision="advance" if i < 8 else "pending")
            db.add(ev)
            db.flush()
            if i < 8:
                stage = stages[i]
                rec = Recruitment(candidate_id=c.id, job_id=j.id, basis_evaluation_id=ev.id, stage=stage)
                db.add(rec)
                db.flush()
                workflow = ["pending_interview", "group_interview", "technical_1", "technical_2", "hr", "intent", "offer_signed", "onboarded"]
                reached = workflow[:workflow.index(stage)+1] if stage in workflow else ["pending_interview", "group_interview", stage]
                previous = ""
                for s in reached:
                    db.add(StageEvent(recruitment_id=rec.id, from_stage=previous, to_stage=s))
                    previous = s
        from docx import Document
        document = Document()
        document.add_paragraph("合成示例简历，仅用于界面预览。")
        document.add_paragraph(text)
        document.save(Path(app.config["DATA_DIR"]) / "uploads" / "example.docx")

if __name__ == "__main__":
    from waitress import serve
    print("Synthetic preview: http://127.0.0.1:5056 | demo / Preview-only-2026", flush=True)
    serve(app, host="127.0.0.1", port=5056, threads=16)

