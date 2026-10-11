import io
import json
from copy import deepcopy
import pytest
from docx import Document
from sqlalchemy import select, func
from app.db import transaction
from app.models import Resume, Candidate, Evaluation, Recruitment, Task, TaskEvent, StageEvent
from app.tasks import execute_task
from app.ai import validate_report


CRITERIA = [
    {"key": "python", "type": "must", "name": "Python", "requirement": "掌握Python", "max_score": 70, "scoring_rule": "有使用依据70，否则0"},
    {"key": "internship", "type": "bonus", "name": "实习", "requirement": "相关实习", "max_score": 30, "scoring_rule": "有实习30，否则0"},
]
RAW = {"relevant_information": {"skills": ["Python"], "internships": [], "projects": ["Python项目"]},
       "criteria_results": [{"criterion_key": "python", "status": "partial", "score": 50, "reason": "仅项目经验",
                              "evidence": [{"quote": "使用Python完成项目", "location": "项目经历"}]},
                             {"criterion_key": "internship", "status": "unknown", "score": 0, "reason": "未写实习", "evidence": []}],
       "strengths": ["项目实践"], "gaps": ["实习待核实"], "interview_questions": ["介绍Python项目"]}


def job(client, headers, name="后端实习生"):
    response = client.post("/api/jobs", json={"name": name, "jd": "Python后端开发实习生", "criteria": deepcopy(CRITERIA)}, headers=headers)
    assert response.status_code == 201, response.json
    return response.json


def ready_resume(app):
    with transaction(app, write=True) as db:
        candidate = Candidate(name="测试候选人", phone="13800000000")
        db.add(candidate)
        db.flush()
        resume = Resume(candidate_id=candidate.id, filename="test.docx", path="test.docx", sha256="abc", status="succeeded", markdown="测试候选人\n使用Python完成项目\n13800000000")
        db.add(resume)
        db.flush()
        return resume.id, candidate.id


def create(client, headers, resume_id, jobs, key="batch-1"):
    response = client.post("/api/evaluations", json={"resume_id": resume_id, "job_ids": [j["id"] for j in jobs], "idempotency_key": key}, headers=headers)
    assert response.status_code in (200, 202), response.json
    return response.json["items"]


def fake_call(app, schema, instructions, payload, on_delta=None):
    if on_delta:
        text = json.dumps(RAW, ensure_ascii=False)
        on_delta(text[:100])
        on_delta(text[100:])
    return deepcopy(RAW), {"model": "test-model", "response_id": "resp_test", "usage": {"total_tokens": 100}}


def test_auth_csrf_setup(app, client):
    anonymous = app.test_client()
    assert anonymous.get("/api/jobs").status_code == 401
    assert client.post("/api/jobs", json={}).status_code == 403
    assert client.get("/setup").status_code == 302
    for page in ["/", "/jobs", "/upload", "/evaluations", "/candidates", "/settings"]:
        assert client.get(page).status_code == 200


def test_score_constraints(client, headers):
    bad = deepcopy(CRITERIA)
    bad[0]["max_score"] = 69
    assert client.post("/api/jobs", json={"name": "test", "jd": "jd", "criteria": bad}, headers=headers).status_code == 400
    bad[0]["max_score"] = 70
    bad[1]["key"] = "python"
    assert client.post("/api/jobs", json={"name": "test", "jd": "jd", "criteria": bad}, headers=headers).status_code == 400


def test_multi_job_snapshot_stream_and_manual_advance(app, client, headers, monkeypatch):
    monkeypatch.setattr("app.tasks.structured_call", fake_call)
    a, b = job(client, headers), job(client, headers, "数据岗位")
    rid, cid = ready_resume(app)
    items = create(client, headers, rid, [a, b])
    assert create(client, headers, rid, [a, b]) == items
    changed = deepcopy(CRITERIA)
    changed[0]["name"] = "新版Python"
    assert client.patch(f'/api/jobs/{a["id"]}', json={"version": a["version"], "criteria": changed}, headers=headers).status_code == 200
    for item in items:
        execute_task(app, item["task_id"])
    data = client.get(f'/api/evaluations/{items[0]["evaluation_id"]}').json
    assert data["status"] == "succeeded"
    assert data["total_score"] == 50
    assert data["snapshot"]["criteria"][0]["name"] == "Python"
    assert "13800000000" not in data["snapshot"]["markdown"]
    assert data["report"]["hard_requirement_failures"][0]["name"] == "Python"
    stream = client.get(f'/api/tasks/{items[0]["task_id"]}/events').data.decode()
    assert "event: completed" in stream
    with transaction(app) as db:
        final_event = db.scalar(select(TaskEvent).where(TaskEvent.task_id == items[0]["task_id"], TaskEvent.type == "completed"))
    replay = client.get(f'/api/tasks/{items[0]["task_id"]}/events', headers={"Last-Event-ID": str(final_event.id - 1)}).data.decode()
    assert "event: completed" in replay
    for _ in range(2):
        assert client.put(f'/api/evaluations/{items[0]["evaluation_id"]}/decision', json={"decision": "advance"}, headers=headers).status_code == 200
    assert client.put(f'/api/evaluations/{items[1]["evaluation_id"]}/decision', json={"decision": "reject"}, headers=headers).status_code == 200
    with transaction(app) as db:
        assert db.scalar(select(func.count()).select_from(Recruitment)) == 1
    stats = client.get("/api/dashboard").json
    assert stats["stats"]["evaluated"] == 1
    assert stats["funnel"][0]["count"] == 2
    assert stats["stats"]["advanced"] == 1


def test_hard_unknown_and_bad_evidence():
    snapshot = {"criteria": deepcopy(CRITERIA), "markdown": "使用Python完成项目", "candidate": {}, "job_name": "test"}
    snapshot["criteria"][1]["type"] = "must"
    report = validate_report(deepcopy(RAW), snapshot)
    assert len(report["hard_requirements_to_verify"]) == 1
    assert len(report["hard_requirement_failures"]) == 1
    raw = deepcopy(RAW)
    raw["criteria_results"][0]["evidence"][0]["quote"] = "编造经历"
    with pytest.raises(ValueError, match="引用"):
        validate_report(raw, snapshot)
    raw = deepcopy(RAW)
    raw["criteria_results"][0]["score"] = 71
    with pytest.raises(ValueError, match="范围"):
        validate_report(raw, snapshot)


def test_stage_skip_correction_and_retraction(app, client, headers, monkeypatch):
    monkeypatch.setattr("app.tasks.structured_call", fake_call)
    j = job(client, headers)
    rid, _ = ready_resume(app)
    item = create(client, headers, rid, [j])[0]
    execute_task(app, item["task_id"])
    eid = item["evaluation_id"]
    rec = client.put(f"/api/evaluations/{eid}/decision", json={"decision": "advance"}, headers=headers).json["recruitment"]
    result = client.post(f'/api/recruitments/{rec["id"]}/stages', json={"stage": "technical_2", "version": rec["version"]}, headers=headers)
    assert result.status_code == 200
    rec = result.json
    assert client.post(f'/api/recruitments/{rec["id"]}/stages', json={"stage": "hr", "version": 1}, headers=headers).status_code == 409
    counts = client.get("/api/dashboard").json["funnel"]
    assert next(r["count"] for r in counts if r["label"] == "技术二面") == 1
    assert next(r["count"] for r in counts if r["label"] == "技术一面") == 0
    rec = client.post(f'/api/recruitments/{rec["id"]}/corrections', json={"event_id": rec["events"][-1]["id"], "stage": "group_interview", "version": rec["version"]}, headers=headers).json
    assert next(r["count"] for r in client.get("/api/dashboard").json["funnel"] if r["label"] == "技术二面") == 0
    client.put(f"/api/evaluations/{eid}/decision", json={"decision": "reject"}, headers=headers)
    assert client.get("/api/recruitments").json["total"] == 0
    client.put(f"/api/evaluations/{eid}/decision", json={"decision": "advance"}, headers=headers)
    assert client.get("/api/recruitments").json["items"][0]["stage"] == "group_interview"


def test_docx_upload_parse_link_delete(app, client, headers):
    doc = Document()
    doc.add_paragraph("张三 本科 软件工程")
    doc.add_paragraph("使用Python完成项目，有完整实践经历。")
    content = io.BytesIO()
    doc.save(content)
    content.seek(0)
    result = client.post("/api/resumes", data={"files": (content, "简历.docx")}, headers=headers)
    assert result.status_code == 202
    item = result.json["items"][0]
    execute_task(app, item["task_id"])
    data = client.get(f'/api/resumes/{item["resume_id"]}').json
    assert data["status"] == "succeeded"
    assert "使用Python" in data["markdown"]
    assert client.post(f'/api/resumes/{item["resume_id"]}/candidate', json={"name": "张三"}, headers=headers).status_code == 200
    assert client.patch(f'/api/resumes/{item["resume_id"]}', json={"markdown": "修改后的内容", "revision": 99}, headers=headers).status_code == 409
    result = client.delete(f'/api/resumes/{item["resume_id"]}', json={"confirm": True}, headers=headers)
    assert result.status_code == 200
    assert client.get(f'/api/resumes/{item["resume_id"]}').status_code == 404
    with transaction(app) as db:
        cleanup = db.scalar(select(Task).where(Task.kind == "cleanup"))
    execute_task(app, cleanup.id)
    assert client.get("/api/maintenance").json["items"] == []


def test_settings_encrypted(client, headers, app):
    from app.models import Setting
    response = client.put("/api/settings/model", json={"model": "test-model", "base_url": "https://api.openai.com/v1", "api_key": "secret-testing-key"}, headers=headers)
    assert response.status_code == 200
    assert "secret-testing-key" not in client.get("/api/settings/model").data.decode()
    with transaction(app) as db:
        assert db.get(Setting, "api_key").value != "secret-testing-key"


def test_task_failure_does_not_expose_provider_input(app, client, headers, monkeypatch):
    def fail(*args, **kwargs):
        raise RuntimeError("sensitive-resume-content secret-key")
    monkeypatch.setattr("app.tasks.structured_call", fail)
    j = job(client, headers)
    rid, _ = ready_resume(app)
    item = create(client, headers, rid, [j])[0]
    execute_task(app, item["task_id"])
    result = client.get(f'/api/evaluations/{item["evaluation_id"]}').json
    assert result["status"] == "failed"
    assert "secret-key" not in result["error"]
    assert client.put(f'/api/evaluations/{item["evaluation_id"]}/decision', json={"decision": "advance"}, headers=headers).status_code == 400


def test_delete_candidate_with_reports_and_history(app, client, headers, monkeypatch):
    monkeypatch.setattr("app.tasks.structured_call", fake_call)
    j = job(client, headers)
    rid, cid = ready_resume(app)
    item = create(client, headers, rid, [j])[0]
    execute_task(app, item["task_id"])
    client.put(f'/api/evaluations/{item["evaluation_id"]}/decision', json={"decision": "advance"}, headers=headers)
    assert client.get(f"/api/deletion-preview?kind=candidate&id={cid}").json["recruitments"] == 1
    response = client.delete(f"/api/candidates/{cid}", json={"confirm": True}, headers=headers)
    assert response.status_code == 200, response.json
    assert client.get("/api/evaluations").json["total"] == 0
    assert client.get("/api/recruitments").json["total"] == 0


def test_background_worker_runs_persistent_task(app, client, headers, monkeypatch):
    import time
    from app.tasks import Worker
    monkeypatch.setattr("app.tasks.structured_call", fake_call)
    j = job(client, headers)
    rid, _ = ready_resume(app)
    item = create(client, headers, rid, [j])[0]
    worker = Worker(app)
    worker.start()
    try:
        deadline = time.monotonic() + 5
        while time.monotonic() < deadline:
            result = client.get(f'/api/tasks/{item["task_id"]}').json
            if result["status"] == "succeeded":
                break
            time.sleep(.1)
        assert result["status"] == "succeeded", result
    finally:
        worker.stop()


def test_restart_marks_incomplete_evaluation_failed(app, client, headers):
    from app.tasks import Worker
    j = job(client, headers)
    rid, _ = ready_resume(app)
    item = create(client, headers, rid, [j])[0]
    with transaction(app, write=True) as db:
        db.get(Task, item["task_id"]).status = "running"
        db.get(Evaluation, item["evaluation_id"]).status = "running"
    worker = Worker(app)
    worker.start()
    worker.stop()
    assert client.get(f'/api/evaluations/{item["evaluation_id"]}').json["status"] == "failed"


def test_upload_type_and_dashboard_date_validation(client, headers):
    response = client.post("/api/resumes", data={"files": (io.BytesIO(b"not a pdf"), "fake.pdf")}, headers=headers)
    assert response.status_code == 400
    assert client.get("/api/dashboard?start=2026-10-12&end=2026-10-11").status_code == 400
    assert client.get("/api/dashboard?start=bad-date").status_code == 400
